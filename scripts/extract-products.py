"""Read Products.xlsx into a traceable, raw import payload.

Only the source columns that have a direct products-table mapping are used by
the importer. The complete source row is retained so columns that do not fit
the first-release schema (prices, stock and accounting fields) are not lost.
"""

import hashlib
import json
from pathlib import Path

import openpyxl


SOURCE = Path("data-template/Products.xlsx")
TARGET = Path(".local/products-import.json")
SHEET_NAME = "Product"
EXPECTED_HEADERS = [
    "เลขสินค้า",
    "สถานะ",
    "ประเภท",
    "ชื่อ",
    "รายละเอียด",
    "ราคาขาย",
    "ต้นทุนเฉลี่ย",
    "คงเหลือ",
    "หมวดหมู่",
    "Barcode",
    "หน่วยนับ",
    "ราคาซื้อ",
    "สถานะสต็อก",
    "บัญชีขายสินค้า",
    "บัญชีขายสินค้าต่างประเทศ",
    "บัญชีสินค้า",
    "บัญชีรับคืนสินค้า",
    "บัญชีส่งคืนสินค้า",
    "บัญชีต้นทุนขาย",
    "บัญชีค่าใช้จ่าย",
]
MAPPING = {
    "product_code": "เลขสินค้า",
    "status": "สถานะ",
    "source_type": "ประเภท",
    "name": "ชื่อ",
    "description": "รายละเอียด",
    "category": "หมวดหมู่",
    "unit": "หน่วยนับ",
}


def main() -> None:
    if not SOURCE.is_file():
        raise SystemExit(f"Source workbook not found: {SOURCE}")

    workbook = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        raise SystemExit(f"Expected sheet {SHEET_NAME!r}; found {workbook.sheetnames!r}")

    sheet = workbook[SHEET_NAME]
    rows = list(sheet.values)
    if not rows:
        raise SystemExit(f"Sheet {SHEET_NAME!r} is empty")

    headers = list(rows[0])
    if headers != EXPECTED_HEADERS:
        raise SystemExit(
            "Unexpected Product headers; refusing to guess a column mapping: "
            f"{headers!r}"
        )

    payload = {
        "schema_version": "products-import-v1",
        "filename": SOURCE.name,
        "sha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
        "sheet": sheet.title,
        "headers": headers,
        "mapping": MAPPING,
        "rows": [
            {"row": row_number, "values": dict(zip(headers, values))}
            for row_number, values in enumerate(rows[1:], start=2)
            if any(value is not None and str(value).strip() for value in values)
        ],
    }

    TARGET.parent.mkdir(exist_ok=True)
    TARGET.write_text(
        json.dumps(payload, ensure_ascii=False, default=str), encoding="utf-8"
    )
    print(
        f"Prepared {len(payload['rows'])} product rows from {sheet.title!r}. "
        "Source workbook unchanged."
    )


if __name__ == "__main__":
    main()
